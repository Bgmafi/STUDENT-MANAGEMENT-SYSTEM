"""
Student Result Management System
---------------------------------
1. Add Student   -> Name, Roll No, Class, 5 subject marks -> Save button
2. Get Result     -> Enter Roll No -> shows Name, Roll No, Class, Total, %, Pass/Fail
3. Show All Result-> Table of every student's result

All data is saved directly into an Excel file (students.xlsx) using openpyxl.
Pass criteria: percentage >= 60
"""

import tkinter as tk
from tkinter import ttk, messagebox
import openpyxl
import os

FILE_NAME = "students.xlsx"
PASS_PERCENT = 60


# ---------------------------------------------------------
# Excel setup
# ---------------------------------------------------------
def init_excel():
    """Create the Excel file with headers if it doesn't already exist."""
    if not os.path.exists(FILE_NAME):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Students"
        ws.append([
            "Name", "Roll No", "Class",
            "Math", "Social Science", "General Science", "English", "Marathi",
            "Total", "Percentage", "Result"
        ])
        wb.save(FILE_NAME)


# ---------------------------------------------------------
# 1. Add Student
# ---------------------------------------------------------
def add_student_window():
    win = tk.Toplevel(root)
    win.title("Add Student")
    win.geometry("300x420")
    win.resizable(False, False)

    fields = ["Name", "Roll No", "Class",
              "Math (out of 100)", "Social Science (out of 100)",
              "General Science (out of 100)", "English (out of 100)",
              "Marathi (out of 100)"]
    entries = []

    for i, label in enumerate(fields):
        tk.Label(win, text=label).grid(row=i, column=0, padx=10, pady=6, sticky="w")
        entry = tk.Entry(win, width=18)
        entry.grid(row=i, column=1, padx=10, pady=6)
        entries.append(entry)

    def save_student():
        name = entries[0].get().strip()
        roll = entries[1].get().strip()
        student_class = entries[2].get().strip()

        if not name or not roll or not student_class:
            messagebox.showerror("Missing Data", "Name, Roll No and Class are required.")
            return

        try:
            marks = [int(entries[i].get()) for i in range(3, 8)]
        except ValueError:
            messagebox.showerror("Invalid Marks", "Marks must be whole numbers (0-100).")
            return

        if any(m < 0 or m > 100 for m in marks):
            messagebox.showerror("Invalid Marks", "Each subject's marks must be between 0 and 100.")
            return

        # check roll no doesn't already exist
        wb = openpyxl.load_workbook(FILE_NAME)
        ws = wb["Students"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if str(row[1]) == roll:
                messagebox.showerror("Duplicate Roll No", "A student with this Roll No already exists.")
                return

        total = sum(marks)
        percentage = total / 5  # 5 subjects, each out of 100
        result = "Pass" if percentage >= PASS_PERCENT else "Fail"

        ws.append([name, roll, student_class, *marks, total, round(percentage, 2), result])
        wb.save(FILE_NAME)

        messagebox.showinfo("Saved", f"{name} saved successfully!\nResult: {result}")
        win.destroy()

    tk.Button(win, text="Save", width=15, bg="#8BC34A",
              command=save_student).grid(row=8, column=0, columnspan=2, pady=25)


# ---------------------------------------------------------
# 2. Get Result
# ---------------------------------------------------------
def get_result_window():
    win = tk.Toplevel(root)
    win.title("Get Result")
    win.geometry("320x320")
    win.resizable(False, False)

    tk.Label(win, text="Enter Roll No:", font=("Arial", 11)).pack(pady=(15, 5))
    roll_entry = tk.Entry(win, width=20)
    roll_entry.pack()

    result_label = tk.Label(win, text="", justify="left", font=("Arial", 11), anchor="w")
    result_label.pack(pady=20, padx=15, fill="x")

    def search_result():
        roll = roll_entry.get().strip()
        if not roll:
            messagebox.showerror("Error", "Please enter a Roll No.")
            return

        wb = openpyxl.load_workbook(FILE_NAME)
        ws = wb["Students"]

        for row in ws.iter_rows(min_row=2, values_only=True):
            if str(row[1]) == roll:
                name, roll_no, student_class = row[0], row[1], row[2]
                total, percentage, result = row[8], row[9], row[10]
                result_label.config(
                    text=(f"Name        : {name}\n"
                          f"Roll No     : {roll_no}\n"
                          f"Class       : {student_class}\n"
                          f"Total Marks : {total} / 500\n"
                          f"Percentage  : {percentage}%\n"
                          f"Result      : {result}")
                )
                return

        result_label.config(text="")
        messagebox.showerror("Not Found", "No student found with this Roll No.")

    tk.Button(win, text="Search", width=15, bg="#03A9F4",
              command=search_result).pack(pady=5)


# ---------------------------------------------------------
# 3. Show All Results
# ---------------------------------------------------------
def show_all_window():
    win = tk.Toplevel(root)
    win.title("All Student Results")
    win.geometry("760x400")

    columns = ("Name", "Roll No", "Class", "Total", "Percentage", "Result")
    tree = ttk.Treeview(win, columns=columns, show="headings")
    for col in columns:
        tree.heading(col, text=col)
        tree.column(col, width=115, anchor="center")
    tree.pack(fill="both", expand=True, padx=10, pady=10)

    wb = openpyxl.load_workbook(FILE_NAME)
    ws = wb["Students"]

    for row in ws.iter_rows(min_row=2, values_only=True):
        name, roll, student_class = row[0], row[1], row[2]
        total, percentage, result = row[8], row[9], row[10]
        tag = "pass" if result == "Pass" else "fail"
        tree.insert("", "end", values=(name, roll, student_class, total, f"{percentage}%", result), tags=(tag,))

    tree.tag_configure("pass", foreground="green")
    tree.tag_configure("fail", foreground="red")


# ---------------------------------------------------------
# Main Menu
# ---------------------------------------------------------
init_excel()

root = tk.Tk()
root.title("Student Result Management System")
root.geometry("360x300")
root.resizable(False, False)

tk.Label(root, text="Student Result System", font=("Arial", 15, "bold")).pack(pady=25)

tk.Button(root, text="1. Add Student", width=28, height=2,
          command=add_student_window).pack(pady=8)
tk.Button(root, text="2. Get Result", width=28, height=2,
          command=get_result_window).pack(pady=8)
tk.Button(root, text="3. Show All Results", width=28, height=2,
          command=show_all_window).pack(pady=8)

root.mainloop()
